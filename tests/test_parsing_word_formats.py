# --- Owner: (assign, see docs/TEAM_SPLIT.md) | tests for Word auto-numbering and ALL-CAPS lead-in clauses ---
# AI-assisted: drafted with Claude (Oct 15 2026). The owner must review it and be able to explain it (IS883 §9).
"""Regression tests for the bug where a lease using Word's automatic numbering came back as one clause.

Common lease templates (for example the Massachusetts Standard Residential Lease Agreement template)
number clauses with Word's numbering feature, so "1.", "2.", ... are not part of the paragraph text,
and they put no blank line between clauses.
"""
import io

import docx

from leasecheck.parsing import load_clauses

HEADINGS = ["THE PARTIES", "RENT", "LATE FEE", "SECURITY DEPOSIT", "RIGHT OF ENTRY", "ATTORNEYS' FEES"]


def _template_docx(numbered: bool, with_subitems: bool = False) -> bytes:
    d = docx.Document()
    d.add_paragraph("MASSACHUSETTS RESIDENTIAL LEASE AGREEMENT")
    for heading in HEADINGS:
        p = d.add_paragraph(style="List Number") if numbered else d.add_paragraph()
        p.add_run(heading + ".").bold = True
        p.add_run(" The Tenant and the Landlord agree to the terms about the " + heading.lower()
                  + " for the Premises, including rent of $2,000 per month.")
        if with_subitems and heading == "RENT":
            for item in ("The Tenant pays by check.", "The Tenant pays on the first day of each month."):
                d.add_paragraph(item, style="List Number 2")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_word_auto_numbering_splits_into_clauses():
    clauses = load_clauses(_template_docx(numbered=True), "template.docx")
    assert len(clauses) == len(HEADINGS)
    assert clauses[0].text.startswith("MASSACHUSETTS RESIDENTIAL LEASE AGREEMENT 1. THE PARTIES.")
    assert clauses[3].text.startswith("4. SECURITY DEPOSIT.")
    assert clauses[3].display_name == "Clause 4"


def test_all_caps_lead_in_splits_without_numbering():
    clauses = load_clauses(_template_docx(numbered=False), "template.docx")
    assert len(clauses) == len(HEADINGS)
    assert clauses[2].text.startswith("LATE FEE.")


def test_sub_items_stay_inside_their_clause():
    clauses = load_clauses(_template_docx(numbered=True, with_subitems=True), "template.docx")
    assert len(clauses) == len(HEADINGS)
    rent = clauses[1].text
    assert rent.startswith("2. RENT.") and "pays by check" in rent and "first day of each month" in rent


def test_bullets_stay_inside_their_clause():
    d = docx.Document()
    d.add_paragraph("RESIDENTIAL LEASE between the Landlord and the Tenant for the Premises and the rent.")
    for heading in ("RENT", "DEFAULT", "NOTICES"):
        p = d.add_paragraph(style="List Number")
        p.add_run(heading + ".").bold = True
        p.add_run(" The Tenant agrees to the terms about " + heading.lower() + " in this lease.")
        if heading == "DEFAULT":
            d.add_paragraph("The Tenant does not pay the Rent;", style="List Bullet")
            d.add_paragraph("The Tenant abandons the Premises.", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    clauses = load_clauses(buf.getvalue(), "bullets.docx")
    # The opening paragraph, then clauses 1 to 3. The bullets do not start new clauses.
    assert [c.display_name for c in clauses] == ["Opening section", "Clause 1", "Clause 2", "Clause 3"]
    default = next(c for c in clauses if "DEFAULT" in c.text)
    assert "does not pay the Rent" in default.text and "abandons the Premises" in default.text
