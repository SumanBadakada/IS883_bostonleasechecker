# --- Owner: Suman | does clause splitting survive different lease formats? (no API calls) ---
# AI-assisted: drafted with Claude Code. The owner must review it and be able to explain it (IS883 §9).
"""Render the same labelled leases in several formats and measure how well parsing recovers the clauses.

A wrong label can come from the model or from the parser handing it the wrong text (two clauses glued
together, one clause cut in half, a page footer inside a clause). This script measures the parser part on
its own, deterministically, so a miss in the main eval can be traced to one or the other.

    python -m eval.format_robustness            # all 20 synthetic leases, every format
    python -m eval.format_robustness --limit 5

For each format it reports the share of clauses recovered exactly, and counts of merged clauses, split
clauses and page furniture (headers/footers) leaking into clause text.
"""
from __future__ import annotations

import argparse
import io
import json
import re
from difflib import SequenceMatcher
from pathlib import Path

from eval.build_test_set import build_leases
from leasecheck.parsing import InputRejected, load_document

FOOTER = "MJMJ"  # tenant initials, as on the real lease where a footer split clause 30
HEADER = "RESIDENTIAL LEASE AGREEMENT"


# ---------- renderers: each returns (bytes, filename) ----------

def _lines(lease: dict, numbered: bool = True) -> list[str]:
    body = [f"{n}. {c['text']}" if numbered else c["text"] for n, c in enumerate(lease["clauses"], start=1)]
    return [HEADER, lease["preamble"]] + body


def docx_numbered(lease):
    import docx
    d = docx.Document()
    for line in _lines(lease):
        d.add_paragraph(line)
    return _save_docx(d), "lease.docx"


def docx_auto_numbered(lease):
    """Clause numbers come from Word's numbering feature, so they are not in the paragraph text."""
    import docx
    d = docx.Document()
    d.add_paragraph(HEADER)
    d.add_paragraph(lease["preamble"])
    for c in lease["clauses"]:
        d.add_paragraph(c["text"], style="List Number")
    return _save_docx(d), "lease.docx"


def docx_unnumbered(lease):
    """No numbers at all: clauses can only be found from their ALL-CAPS lead-ins ("LATE PAYMENT. ...")."""
    import docx
    d = docx.Document()
    for line in _lines(lease, numbered=False):
        d.add_paragraph(line)
    return _save_docx(d), "lease.docx"


def _save_docx(d) -> bytes:
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf(lease, columns: int = 1, furniture: bool = False):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph

    width, height = letter
    margin, gap = 72, 24
    col_w = (width - 2 * margin - gap * (columns - 1)) / columns
    frames = [Frame(margin + i * (col_w + gap), margin, col_w, height - 2 * margin, id=f"c{i}") for i in range(columns)]

    def furniture_on_page(canvas, doc):
        if furniture:
            canvas.setFont("Helvetica", 8)
            canvas.drawString(margin, height - 40, f"{HEADER} - Page {doc.page}")
            canvas.drawString(margin, 36, f"{FOOTER} {doc.page}")

    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=letter, pageTemplates=[PageTemplate(frames=frames, onPage=furniture_on_page)])
    # A large font so clauses run across page (and column) breaks, which is where splitting goes wrong.
    style = ParagraphStyle("body", fontName="Helvetica", fontSize=13, leading=19, spaceAfter=9)
    doc.build([Paragraph(line.replace("&", "&amp;"), style) for line in _lines(lease)])
    return buf.getvalue(), "lease.pdf"


def pdf_text(lease):
    return _pdf(lease)


def pdf_with_header_footer(lease):
    return _pdf(lease, furniture=True)


def pdf_two_column(lease):
    return _pdf(lease, columns=2)


FORMATS = {
    "docx_numbered": docx_numbered,
    "docx_auto_numbered": docx_auto_numbered,
    "docx_unnumbered": docx_unnumbered,
    "pdf_text": pdf_text,
    "pdf_header_footer": pdf_with_header_footer,
    "pdf_two_column": pdf_two_column,
}


# ---------- scoring ----------

def _norm(text: str) -> str:
    text = re.sub(r"^\s*\d{1,2}[.)]\s*", "", text)  # the clause number is not part of the clause
    return re.sub(r"\s+", " ", text).strip().lower()


def score(gold_texts: list[str], predicted: list[str]) -> dict:
    """exact: a predicted clause has (near enough) the same text as the gold clause.
    merged: a predicted clause contains the starts of two or more gold clauses.
    split: a gold clause's start and end land in different predicted clauses.
    leaked: a predicted clause contains page header/footer text."""
    gold = [_norm(t) for t in gold_texts]
    pred = [_norm(t) for t in predicted]
    found = [any(SequenceMatcher(None, g, p).ratio() >= 0.97 for p in pred) for g in gold]
    exact = sum(found)
    merged = sum(sum(g[:40] in p for g in gold) >= 2 for p in pred)
    split = 0
    for g, ok in zip(gold, found):
        if ok:
            continue  # recovered whole, so not split (several clauses can share an opening or closing phrase)
        start = next((i for i, p in enumerate(pred) if g[:40] in p), None)
        end = next((i for i, p in enumerate(pred) if g[-40:] in p), None)
        split += start is not None and end is not None and start != end
    leaked = sum(FOOTER.lower() in p or "- page" in p for p in pred)
    return {"clauses": len(gold), "exact": exact, "merged": merged, "split": split, "leaked_furniture": leaked}


def run(leases: list[dict], formats: dict = FORMATS) -> dict[str, dict]:
    totals: dict[str, dict] = {}
    for name, render in formats.items():
        t = {"clauses": 0, "exact": 0, "merged": 0, "split": 0, "leaked_furniture": 0, "rejected": 0}
        for lease in leases:
            data, filename = render(lease)
            gold = [c["text"] for c in lease["clauses"]]
            try:
                predicted = [c.text for c in load_document(data, filename).clauses]
            except InputRejected:
                t["rejected"] += 1
                t["clauses"] += len(gold)
                continue
            for k, v in score(gold, predicted).items():
                t[k] += v
        t["exact_rate"] = round(t["exact"] / t["clauses"], 3) if t["clauses"] else None
        totals[name] = t
    return totals


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=0)
    args = ap.parse_args()
    leases = build_leases()
    if args.limit:
        leases = leases[: args.limit]
    totals = run(leases)

    print(f"Clause recovery across formats ({len(leases)} leases)\n")
    print(f"{'format':22} {'exact':>7} {'merged':>7} {'split':>6} {'leaked':>7} {'rejected':>9}")
    for name, t in totals.items():
        print(f"{name:22} {t['exact_rate']:>7.1%} {t['merged']:>7} {t['split']:>6} {t['leaked_furniture']:>7} {t['rejected']:>9}")
    out = Path(__file__).resolve().parent / "results"
    out.mkdir(exist_ok=True)
    (out / "format_robustness.json").write_text(json.dumps(totals, indent=2))
    print(f"\nSaved {out / 'format_robustness.json'}")


if __name__ == "__main__":
    main()
