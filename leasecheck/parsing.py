# --- Owner: (assign, see docs/TEAM_SPLIT.md) | reading the upload and splitting it into clauses ---
# AI-assisted: drafted with Claude Code. The owner must review it and be able to explain it (IS883 §9).
"""Turn an uploaded PDF or DOCX into a list of clauses.

No model is used here: text extraction, scanned-PDF detection, the lease check and clause splitting
are plain code, which is cheaper and cannot hallucinate.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

from . import config


class InputRejected(Exception):
    """Raised for inputs we decline. `reason` is a short code, `message` is shown to the user."""

    def __init__(self, reason: str, message: str):
        super().__init__(message)
        self.reason = reason
        self.message = message


@dataclass
class Clause:
    clause_id: int  # our own position counter, used to match the model's answers to clauses
    text: str

    @property
    def display_name(self) -> str:
        """What the tenant sees: the lease's own number ("Clause 3") when it has one."""
        match = re.match(r"\s*(?:section|article|clause|§)?\s*(\d{1,2}(?:\.\d{1,2})*)\s*[.):\-]\s", self.text, re.IGNORECASE)
        if match:
            return f"Clause {match.group(1)}"
        return "Opening section" if self.clause_id == 1 else f"Part {self.clause_id}"


# A PDF averaging fewer extractable characters per page than this is treated as scanned.
MIN_CHARS_PER_PAGE = 80

# Words that almost every residential lease uses. A document needs several of them to count as a lease.
LEASE_TERMS = (
    "landlord", "lessor", "tenant", "lessee", "rent", "premises", "lease",
    "security deposit", "term of", "apartment", "occupan",
)
MIN_LEASE_TERMS = 4


def extract_text(data: bytes, filename: str) -> str:
    """Extract plain text. Raises InputRejected for empty, unsupported, oversized or scanned files."""
    if not data:
        raise InputRejected("empty", "The file is empty. Please upload your lease as a PDF or Word (.docx) file.")
    if len(data) > config.MAX_FILE_MB * 1024 * 1024:
        raise InputRejected("too_large", f"The file is larger than {config.MAX_FILE_MB} MB.")

    name = filename.lower()
    if name.endswith(".pdf"):
        text = _pdf_text(data)
    elif name.endswith(".docx"):
        text = _docx_text(data)
    elif name.endswith(".txt"):  # used by the eval harness and tests
        text = data.decode("utf-8", errors="replace")
    else:
        raise InputRejected("unsupported", "Please upload a PDF or a Word (.docx) file.")

    text = _normalise(text)
    if not text.strip():
        raise InputRejected("empty", "We could not find any text in this file.")
    return text


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, ValueError, KeyError) as exc:
        raise InputRejected("unreadable", "This PDF could not be read. It may be damaged or password-protected.") from exc
    if not pages:
        raise InputRejected("empty", "This PDF has no pages.")
    chars = sum(len(p.strip()) for p in pages)
    if chars / len(pages) < MIN_CHARS_PER_PAGE:
        raise InputRejected(
            "scanned",
            "This looks like a scanned PDF (an image of the pages), so there is no text for us to check. "
            "Please upload the lease as a text-based PDF or a Word file.",
        )
    return "\n".join(pages)


def _docx_text(data: bytes) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise InputRejected("unreadable", "This Word file could not be read.") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def check_is_lease(text: str) -> None:
    """Reject documents that are clearly not leases, before spending any model calls on them."""
    lower = text.lower()
    found = [term for term in LEASE_TERMS if term in lower]
    if len(found) < MIN_LEASE_TERMS:
        raise InputRejected(
            "not_a_lease",
            "This does not look like a residential lease, so we have not checked it. "
            "Please upload the lease agreement itself.",
        )


# A line that starts a new clause: "1.", "12)", "3.1", "(a)" is NOT a new clause (sub-item),
# "Section 4", "ARTICLE V", or a short ALL-CAPS heading such as "SECURITY DEPOSIT".
_NUMBERED = re.compile(r"^\s*(?:section|article|clause|§)?\s*(\d{1,2}|[IVXL]{1,6})(?:\.\d{1,2})*\s*[.):\-]\s+\S", re.IGNORECASE)
_HEADING = re.compile(r"^\s*[A-Z][A-Z0-9 ,'&/\-]{3,60}:?\s*$")


def split_clauses(text: str) -> list[Clause]:
    """Split lease text into clauses using numbering and headings, falling back to paragraphs."""
    lines = [line.rstrip() for line in text.split("\n")]
    starts = [i for i, line in enumerate(lines) if _NUMBERED.match(line) or _HEADING.match(line)]

    blocks: list[str] = []
    if len(starts) >= 3:
        # Anything before the first numbered line is the preamble (parties, address), kept as its own block.
        if starts[0] > 0:
            blocks.append("\n".join(lines[: starts[0]]))
        for a, b in zip(starts, starts[1:] + [len(lines)]):
            blocks.append("\n".join(lines[a:b]))
        blocks = _merge_headings(blocks)
    else:
        blocks = [b for b in re.split(r"\n\s*\n", text)]

    cleaned = [re.sub(r"\s*\n\s*", " ", b).strip() for b in blocks]
    cleaned = [b for b in cleaned if len(b) >= 25]
    if len(cleaned) > config.MAX_CLAUSES:
        raise InputRejected(
            "too_long",
            f"This document splits into {len(cleaned)} clauses, more than the {config.MAX_CLAUSES} we can check at once.",
        )
    return [Clause(clause_id=i + 1, text=b) for i, b in enumerate(cleaned)]


def _merge_headings(blocks: list[str]) -> list[str]:
    """A bare heading line ("SECURITY DEPOSIT") belongs with the text after it."""
    merged: list[str] = []
    carry = ""
    for block in blocks:
        if _HEADING.match(block.strip()) and "\n" not in block.strip():
            carry = (carry + " " + block.strip()).strip()
            continue
        merged.append((carry + "\n" + block).strip() if carry else block)
        carry = ""
    if carry:
        merged.append(carry)
    return merged


def load_clauses(data: bytes, filename: str) -> list[Clause]:
    """The full deterministic front end: extract, check it is a lease, split."""
    text = extract_text(data, filename)
    check_is_lease(text)
    clauses = split_clauses(text)
    if not clauses:
        raise InputRejected("empty", "We could not find any clauses in this document.")
    return clauses
